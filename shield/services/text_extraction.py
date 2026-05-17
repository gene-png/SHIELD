"""Best-effort text extraction from common office formats.

Used at upload time by the repository writer to populate
`Artifact.body_text` so downstream AI processing has structured-ish text
to work with without re-parsing the binary on every call.

Hard rule: extraction is **best effort**. Failure to extract is NOT
fatal — the file is still saved to the human lane with the binary intact;
`body_text` just remains None. The AI processing layer must handle that
gracefully (see p1_techdebt/routes.py `extract` route).

Supported formats:
- PDF (`application/pdf`, `*.pdf`)        via pypdf
- DOCX (Word 2007+, `*.docx`)             via python-docx
- XLSX (Excel 2007+, `*.xlsx`)            via openpyxl
- Plain text (`text/*`, `*.txt|md|csv`)   direct decode

Unsupported binary formats (legacy `.doc`, `.xls`, images, etc.) return
None; the caller should leave body_text unset for those.
"""
from __future__ import annotations

import logging
from typing import BinaryIO

import openpyxl
import pypdf
from docx import Document

log = logging.getLogger(__name__)

# Mime-type fragments → extractor name. Checked before extension fallback.
_MIME_HINTS = {
    "pdf": "_extract_pdf",
    "wordprocessingml": "_extract_docx",
    "spreadsheetml": "_extract_xlsx",
    "text/": "_extract_text",
}

_EXT_HINTS = {
    "pdf": "_extract_pdf",
    "docx": "_extract_docx",
    "xlsx": "_extract_xlsx",
    "txt": "_extract_text",
    "md": "_extract_text",
    "csv": "_extract_text",
    "json": "_extract_text",
    "xml": "_extract_text",
    "log": "_extract_text",
}


def extract_text(
    stream: BinaryIO,
    *,
    mime_type: str | None,
    filename: str | None,
) -> str | None:
    """Dispatch on mime_type then filename extension. Returns None if unsupported.

    Errors during extraction are caught and logged; the function returns
    None on failure so the caller can fall through to "no body_text" rather
    than aborting the upload.
    """
    mt = (mime_type or "").lower()
    ext = ""
    if filename and "." in filename:
        ext = filename.rsplit(".", 1)[-1].lower()

    extractor_name: str | None = None
    for hint, name in _MIME_HINTS.items():
        if hint in mt:
            extractor_name = name
            break
    if extractor_name is None and ext in _EXT_HINTS:
        extractor_name = _EXT_HINTS[ext]

    if extractor_name is None:
        return None

    try:
        return globals()[extractor_name](stream)
    except Exception as e:
        log.warning("text extraction failed for filename=%r mime=%r: %s", filename, mime_type, e)
        return None


def _extract_pdf(stream: BinaryIO) -> str:
    stream.seek(0)
    reader = pypdf.PdfReader(stream)
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text:
            parts.append(text)
    return "\n\n".join(parts).strip()


def _extract_docx(stream: BinaryIO) -> str:
    stream.seek(0)
    doc = Document(stream)
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join((cell.text or "").strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)
    return "\n".join(parts).strip()


def _extract_xlsx(stream: BinaryIO) -> str:
    stream.seek(0)
    wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts).strip()


def _extract_text(stream: BinaryIO) -> str:
    stream.seek(0)
    raw = stream.read()
    if isinstance(raw, str):
        return raw.strip()
    # UTF-8 with replacement is robust enough for client uploads; charset
    # detection is a v1.0 follow-up if it becomes a real problem.
    return raw.decode("utf-8", errors="replace").strip()
