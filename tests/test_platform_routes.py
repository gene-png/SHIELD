"""End-to-end-ish route tests for the v0.3-v1.1 work.

These exercise the spec invariants through the HTTP layer, not just the
spine functions:
  - The picker's AI-origin acknowledgment gate fires on relink.
  - Evidence upload attaches to the right QuestionnaireResponse.
  - Submit locks every response and stops further edits.
  - Attribution downgrade rejects upgrades.
  - The role gate redirects CLIENT-role users out of platform URLs.

Tests run against TestConfig (SQLite in-memory + AI_MODE=fixture), so
no Postgres triggers, no Keycloak, no Anthropic calls.
"""
from __future__ import annotations

import io

from shield.extensions import db
from shield.models import (
    Artifact,
    Origin,
    Project,
    QuestionnaireResponse,
    TrustTier,
)

# --------------------------------------------------------------------
# Picker — AI-origin acknowledgment gate on relink
# --------------------------------------------------------------------

def test_relink_to_ai_origin_without_ack_is_rejected(
    admin_client, acme, ai_capability_list, p2_project,
):
    """Relinking to an AI-origin capability list MUST require explicit ack."""
    response = admin_client.post(
        f"/projects/{p2_project.id}/relink-capability-list",
        data={"capability_list_id": ai_capability_list.id},
        follow_redirects=False,
    )
    # Form re-renders with an error (200), not the redirect (302) you'd
    # see on success. The project's capability_list_version_id stays None.
    refreshed = db.session.get(Project, p2_project.id)
    assert refreshed.capability_list_version_id is None
    assert response.status_code == 200


def test_relink_to_ai_origin_with_ack_succeeds(
    admin_client, acme, ai_capability_list, p2_project,
):
    """With the acknowledgment, the relink is allowed."""
    response = admin_client.post(
        f"/projects/{p2_project.id}/relink-capability-list",
        data={
            "capability_list_id": ai_capability_list.id,
            "acknowledged_ai_reuse": "yes",
        },
        follow_redirects=False,
    )
    assert response.status_code == 302  # redirect to workspace
    refreshed = db.session.get(Project, p2_project.id)
    assert refreshed.capability_list_version_id == ai_capability_list.id


# --------------------------------------------------------------------
# P2 evidence upload tied to a control (spec §8.2)
# --------------------------------------------------------------------

def test_evidence_upload_attaches_to_response(admin_client, admin, p2_project):
    """Uploading evidence on a control links the new artifact to the response."""
    # Need an answer first — evidence attaches to an existing response.
    r = QuestionnaireResponse(
        project_id=p2_project.id, framework="cisa_ztmm_v2",
        control_id="ZTMM.IDENT.1", answer="implemented", rationale="",
        trust_tier=TrustTier.ADMIN_ASSISTED, attributed_user_id=admin.id,
    )
    db.session.add(r)
    db.session.commit()

    response = admin_client.post(
        f"/platform/zero-trust/project/{p2_project.id}/evidence/ZTMM.IDENT.1",
        data={"evidence_file": (io.BytesIO(b"some evidence pdf bytes"), "evidence.pdf")},
        content_type="multipart/form-data",
        follow_redirects=False,
    )
    assert response.status_code == 302
    refreshed = db.session.get(QuestionnaireResponse, r.id)
    assert refreshed.evidence_artifact_id is not None
    art = db.session.get(Artifact, refreshed.evidence_artifact_id)
    assert art is not None
    assert art.origin == Origin.HUMAN_INPUT
    assert art.trust_tier == TrustTier.CLIENT_PROVIDED_EVIDENCE
    assert art.stage == "evidence"


# --------------------------------------------------------------------
# P2 submit / lock (spec §10 decision #3)
# --------------------------------------------------------------------

def test_submit_locks_every_response(admin_client, admin, p2_project):
    """Submit must mark every response locked + flip project.stage."""
    for cid in ("ZTMM.IDENT.1", "ZTMM.IDENT.2"):
        db.session.add(QuestionnaireResponse(
            project_id=p2_project.id, framework="cisa_ztmm_v2",
            control_id=cid, answer="implemented", rationale="",
            trust_tier=TrustTier.ADMIN_ASSISTED, attributed_user_id=admin.id,
        ))
    db.session.commit()

    response = admin_client.post(
        f"/platform/zero-trust/project/{p2_project.id}/submit",
        follow_redirects=False,
    )
    assert response.status_code == 302

    db.session.expire_all()
    assert all(
        r.locked for r in db.session.query(QuestionnaireResponse)
                                    .filter_by(project_id=p2_project.id)
    )
    assert db.session.get(Project, p2_project.id).stage == "submitted"


# --------------------------------------------------------------------
# P2 attribution downgrade — never an upgrade (spec §10 decision #3)
# --------------------------------------------------------------------

def test_attribution_can_only_be_downgraded(admin_client, admin, p2_project):
    """ADMIN_ASSISTED -> CLIENT_ASSERTED is an upgrade; must be rejected."""
    r = QuestionnaireResponse(
        project_id=p2_project.id, framework="cisa_ztmm_v2",
        control_id="ZTMM.IDENT.1", answer="implemented", rationale="",
        trust_tier=TrustTier.ADMIN_ASSISTED, attributed_user_id=admin.id,
    )
    db.session.add(r)
    db.session.commit()

    # Try to upgrade — should not change the row.
    admin_client.post(
        f"/platform/zero-trust/project/{p2_project.id}/answer/ZTMM.IDENT.1/downgrade",
        data={"trust_tier": "client_asserted"},
    )
    db.session.expire_all()
    assert db.session.get(QuestionnaireResponse, r.id).trust_tier == TrustTier.ADMIN_ASSISTED

    # Now properly downgrade — must succeed.
    admin_client.post(
        f"/platform/zero-trust/project/{p2_project.id}/answer/ZTMM.IDENT.1/downgrade",
        data={"trust_tier": "admin_entered_on_behalf"},
    )
    db.session.expire_all()
    assert (
        db.session.get(QuestionnaireResponse, r.id).trust_tier
        == TrustTier.ADMIN_ENTERED_ON_BEHALF
    )


# --------------------------------------------------------------------
# Role gate — CLIENT-role users never reach the platforms (spec §6.6)
# --------------------------------------------------------------------

def test_client_role_redirected_from_platform_routes(client_role_client):
    """Any non-intake/auth URL must redirect a CLIENT-role user to /intake."""
    for path in (
        "/platform/tech-debt/",
        "/platform/zero-trust/",
        "/platform/attack-surface/",
        "/clients/",
        "/repository/",
    ):
        r = client_role_client.get(path, follow_redirects=False)
        assert r.status_code == 302, f"{path!r} did not redirect"
        assert r.headers["Location"].endswith("/intake/"), (
            f"{path!r} redirected to {r.headers['Location']!r}, not /intake/"
        )


def test_client_role_can_reach_intake(client_role_client):
    """The intake surface itself is allowed for CLIENT-role users."""
    r = client_role_client.get("/intake/")
    assert r.status_code == 200


# --------------------------------------------------------------------
# Defensive RBAC — REVIEWER can browse, can NOT mutate (v1.3)
# --------------------------------------------------------------------

def test_reviewer_can_browse_platform_indexes(reviewer_client):
    """Reviewer is the audit persona: read access across all platforms."""
    for path in (
        "/platform/tech-debt/",
        "/platform/zero-trust/",
        "/platform/attack-surface/",
        "/clients/",
        "/repository/",
    ):
        r = reviewer_client.get(path)
        assert r.status_code == 200, f"reviewer cannot view {path!r}"


def test_reviewer_blocked_from_mutating_routes(reviewer_client):
    """Every @admin_only route must reject reviewer with a redirect."""
    for path in (
        "/platform/tech-debt/new",
        "/platform/zero-trust/new",
        "/platform/attack-surface/new",
    ):
        r = reviewer_client.get(path, follow_redirects=False)
        assert r.status_code == 302, f"reviewer NOT blocked from {path!r}"


def test_p3_coverage_run_xlsx_export(admin_client, acme, admin):
    """The .xlsx export for a coverage run returns a 4-sheet workbook."""
    import io

    import openpyxl

    from shield.models import (
        CoverageFinding,
        CoverageRun,
        PlatformType,
        Project,
    )
    cl = acme.capability_lists[0]
    project = Project(
        client_id=acme.id, platform=PlatformType.ATTACK_SURFACE,
        name="ATT&CK test", stage="intake",
        capability_list_version_id=cl.id, created_by_id=admin.id,
    )
    db.session.add(project)
    db.session.commit()

    run = CoverageRun(
        project_id=project.id,
        capability_list_version_id=cl.id,
        summary={
            "headline": "You are blind to 2 techniques.",
            "covered": 1, "partial": 1, "uncovered": 2, "total_techniques": 4,
            "top_three_blind_spots": ["T1566", "T1486"],
        },
    )
    db.session.add(run)
    db.session.flush()
    db.session.add(CoverageFinding(
        coverage_run_id=run.id, technique_id="T1566",
        coverage="uncovered", detection_tools=[], prevention_tools=[],
        response_tools=[], rationale="No email gateway.",
    ))
    db.session.add(CoverageFinding(
        coverage_run_id=run.id, technique_id="T1486",
        coverage="partial", detection_tools=["EDR"], prevention_tools=[],
        response_tools=["Backup"], rationale="Backups but no prevention.",
    ))
    db.session.commit()

    r = admin_client.get(
        f"/platform/attack-surface/project/{project.id}/run/{run.id}/export.xlsx"
    )
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["Content-Type"]

    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "Coverage", "Gaps", "Methodology"]
    # Gaps sheet should list both findings (partial + uncovered)
    gap_rows = list(wb["Gaps"].iter_rows(values_only=True))
    assert gap_rows[0][0] == "Technique ID"
    assert {row[0] for row in gap_rows[1:]} == {"T1566", "T1486"}
    # "uncovered" gets re-labeled to "not covered" in the XLSX (the DB
    # value stays "uncovered" so no migration is needed). Column index
    # 3 on the Gaps sheet is "Coverage".
    coverage_labels = {row[3] for row in gap_rows[1:]}
    assert coverage_labels == {"not covered", "partial"}


def test_reviewer_can_view_audit_log(reviewer_client):
    """Audit log is the reviewer's primary read surface."""
    r = reviewer_client.get("/audit/")
    assert r.status_code == 200


def test_capability_list_xlsx_export(admin_client, acme):
    """The .xlsx export returns a valid openpyxl-readable workbook."""
    import io

    import openpyxl

    from shield.models import CapabilityListItem
    cl = acme.capability_lists[0]
    db.session.add(CapabilityListItem(
        capability_list_id=cl.id, name="Splunk", vendor="Splunk Inc",
        category="SIEM", function="log aggregation",
        annual_cost_usd=420000, license_count=1,
    ))
    db.session.commit()

    r = admin_client.get(f"/clients/{acme.id}/capability-list/{cl.id}/export.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["Content-Type"]
    assert 'attachment; filename="' in r.headers["Content-Disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert "Overview" in wb.sheetnames
    assert "Items" in wb.sheetnames
    items_rows = list(wb["Items"].iter_rows(values_only=True))
    assert items_rows[0][0] == "Name"           # header row
    assert items_rows[1][0] == "Splunk"         # body row


def test_client_blocked_from_audit_log(client_role_client):
    """CLIENT role should never see the audit log."""
    r = client_role_client.get("/audit/", follow_redirects=False)
    assert r.status_code == 302
    assert r.headers["Location"].endswith("/intake/")


def test_reviewer_can_promote_ai_artifacts(reviewer_client, admin, p2_project):
    """Reviewer is one of the two roles allowed to promote AI artifacts."""
    from shield.models import Artifact, Origin, ReuseStatus
    from shield.spine.repository import write_ai_artifact
    art = write_ai_artifact(
        project=p2_project, stage="current_state_assessment",
        title="ai", body_text="{}",
        input_artifact_ids=[], prompt_version="p2_posture.v1", model="x",
    )
    r = reviewer_client.post(
        f"/repository/artifact/{art.id}/promote",
        data={"reason": "Reviewed against three controls; signs off."},
        follow_redirects=False,
    )
    assert r.status_code == 302
    refreshed = db.session.get(Artifact, art.id)
    assert refreshed.origin == Origin.AI_GENERATED   # origin unchanged
    assert refreshed.reuse_status == ReuseStatus.APPROVED
